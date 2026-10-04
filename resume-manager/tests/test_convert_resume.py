"""Layout tests for assets/convert_resume.py DOCX export.

Run: python -m unittest discover -s resume-manager/tests
Requires: python-docx
"""

import sys
import tempfile
import unittest
from pathlib import Path

ASSETS = Path(__file__).resolve().parent.parent / "assets"
sys.path.insert(0, str(ASSETS))

import docx  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT  # noqa: E402

import convert_resume  # noqa: E402

RESUME_MD = """# Jane Q. Doe

**Senior Software Engineer**

Chicago, IL | (555) 555-0100 | jane@example.com | [linkedin.com/in/jane](https://linkedin.com/in/jane)

## Professional Summary

Senior software engineer building data platforms with C#/.NET and TypeScript.

## Technical Skills

- **Languages:** C#, TypeScript, SQL

## Professional Experience

### Acme Corp | Jul 2023 – Present
*Lead Software Engineer* | Remote
- Cut onboarding from 3 days to 5 minutes by building an operations console.

## Open Source

### [Widget](https://example.com/widget) | Sep 2026 – Present
- Reduced token use by ~65% by building a prompt slicer.
"""

LEGACY_TITLE_MD = """# Jane Q. Doe

Jane Q. Doe | Chicago, IL | jane@example.com

# Senior Software Engineer

Senior software engineer building data platforms with C#/.NET and TypeScript.

## Technical Skills

- **Languages:** C#
"""


COVER_LETTER_MD = """Jane Q. Doe
Chicago, IL | jane@example.com | (555) 555-0100 | linkedin.com/in/jane

October 3, 2026

Hiring Team
Acme Corp

I build data platforms with C#/.NET and own them in production.
"""


def build(md_text: str, stem: str = "resume") -> docx.document.Document:
    tmp = Path(tempfile.mkdtemp())
    md_path = tmp / f"{stem}.md"
    md_path.write_text(md_text, encoding="utf-8")
    out = tmp / f"{stem}.docx"
    convert_resume.convert_md_to_docx(md_path, out)
    return docx.Document(str(out))


def find(doc, text: str):
    for p in doc.paragraphs:
        if text in p.text:
            return p
    raise AssertionError(f"paragraph containing {text!r} not found")


def max_size(p) -> float:
    return max((r.font.size.pt for r in p.runs if r.font.size), default=0)


class ConvertResumeLayoutTest(unittest.TestCase):
    def test_subtitle_is_smaller_than_name(self):
        doc = build(RESUME_MD)
        self.assertLess(max_size(find(doc, "Senior Software Engineer")), max_size(find(doc, "Jane Q. Doe")))

    def test_later_h1_renders_as_subtitle_not_name(self):
        doc = build(LEGACY_TITLE_MD)
        name = doc.paragraphs[0]
        title = find(doc, "Senior Software Engineer")
        self.assertEqual(name.text, "Jane Q. Doe")
        self.assertLess(max_size(title), max_size(name))

    def test_summary_prose_is_left_aligned(self):
        for md in (RESUME_MD, LEGACY_TITLE_MD):
            p = find(build(md), "building data platforms")
            self.assertNotEqual(p.alignment, WD_ALIGN_PARAGRAPH.CENTER)

    def test_contact_line_is_centered(self):
        p = find(build(RESUME_MD), "jane@example.com")
        self.assertEqual(p.alignment, WD_ALIGN_PARAGRAPH.CENTER)

    def test_company_line_right_aligns_dates_with_tab(self):
        p = find(build(RESUME_MD), "Acme Corp")
        self.assertEqual(p.text, "Acme Corp\tJul 2023 – Present")
        stops = list(p.paragraph_format.tab_stops)
        self.assertTrue(any(s.alignment == WD_TAB_ALIGNMENT.RIGHT for s in stops))
        self.assertTrue(all(r.bold for r in p.runs if r.text.strip()))

    def test_role_line_is_italic_with_right_aligned_location(self):
        p = find(build(RESUME_MD), "Lead Software Engineer")
        self.assertEqual(p.text, "Lead Software Engineer\tRemote")
        self.assertTrue(all(r.italic for r in p.runs if r.text.strip()))

    def test_linked_project_heading_keeps_label_and_date(self):
        p = find(build(RESUME_MD), "Sep 2026")
        self.assertIn("Widget", "".join(t.text for t in p._p.iter(docx.oxml.ns.qn("w:t"))))
        self.assertTrue(p.text.endswith("\tSep 2026 – Present"))

    def test_visible_url_beside_project_name_is_not_bold(self):
        md = RESUME_MD + "\n### Gadget ([example.com/gadget](https://example.com/gadget)) | Feb 2026 – Present\n"
        p = find(build(md), "Gadget")
        qn = docx.oxml.ns.qn
        link_runs = [r for h in p._p.iter(qn("w:hyperlink")) for r in h.iter(qn("w:r"))]
        self.assertTrue(link_runs)
        self.assertTrue(all(r.find(qn("w:rPr")).find(qn("w:b")) is None for r in link_runs))
        self.assertTrue(p.runs[0].bold)

    def test_no_tables_and_black_text(self):
        doc = build(RESUME_MD)
        self.assertEqual(len(doc.tables), 0)
        for p in doc.paragraphs:
            for r in p.runs:
                color = r.font.color.rgb if r.font.color and r.font.color.type else None
                self.assertIn(str(color) if color else None, (None, "000000"), p.text)

    def test_cover_letter_contact_line_is_centered_with_name(self):
        doc = build(COVER_LETTER_MD, stem="cover_letter")
        self.assertEqual(find(doc, "Jane Q. Doe").alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertEqual(find(doc, "jane@example.com").alignment, WD_ALIGN_PARAGRAPH.CENTER)

    def test_cover_letter_body_after_letterhead_is_left_aligned(self):
        doc = build(COVER_LETTER_MD, stem="cover_letter")
        for text in ("October 3, 2026", "Hiring Team", "I build data platforms"):
            self.assertNotEqual(find(doc, text).alignment, WD_ALIGN_PARAGRAPH.CENTER, text)

    def test_cover_letter_consecutive_lines_stay_in_one_block(self):
        doc = build(COVER_LETTER_MD, stem="cover_letter")
        block = find(doc, "Hiring Team")
        self.assertEqual(block.text, "Hiring Team\nAcme Corp")

    def test_bullets_use_compact_indent(self):
        p = find(build(RESUME_MD), "Cut onboarding")
        self.assertLessEqual(p.paragraph_format.left_indent.inches, 0.25)


if __name__ == "__main__":
    unittest.main()
